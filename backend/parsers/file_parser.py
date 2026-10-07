"""Парсер файлов: все форматы (PDF, DOCX, DOC, TXT, MD, JSON, PPTX, XLSX, CSV, HTML, RTF, ODT, EPUB, изображения)."""
import csv
import io
import json
import logging
import re
import subprocess
import shutil
from pathlib import Path

import chardet

from backend.config import OCR_ENABLED
from backend.parsers.ocr import ocr_image, ocr_pdf

logger = logging.getLogger(__name__)


def parse_file(path: str) -> list[dict]:
    """
    Распарсить файл по расширению.
    Возвращает список документов: [{title, text, date?, meta?}]
    """
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"File not found: {path}")

    ext = p.suffix.lower()
    logger.info("Parsing file: %s (ext=%s)", p.name, ext)

    try:
        match ext:
            case ".pdf":
                return _parse_pdf(p)
            case ".docx":
                return _parse_docx(p)
            case ".doc":
                return _parse_doc(p)
            case ".txt":
                return _parse_txt(p)
            case ".md":
                return _parse_md(p)
            case ".json":
                return _parse_json(p)
            case ".pptx":
                return _parse_pptx(p)
            case ".xlsx":
                return _parse_xlsx(p)
            case ".csv":
                return _parse_csv(p)
            case ".html" | ".htm":
                return _parse_html(p)
            case ".rtf":
                return _parse_rtf(p)
            case ".odt":
                return _parse_odt(p)
            case ".epub":
                return _parse_epub(p)
            case ".png" | ".jpg" | ".jpeg" | ".bmp" | ".tiff":
                return _parse_image(p)
            case _:
                raise ValueError(f"Unsupported file format: {ext}")
    except Exception as e:
        logger.error("Failed to parse %s: %s", path, e)
        raise


def _parse_pdf(p: Path) -> list[dict]:
    """PDF: pdfplumber для текста, OCR для сканов."""
    import pdfplumber

    docs: list[dict] = []
    full_text_parts: list[str] = []
    needs_ocr = False

    with pdfplumber.open(p) as pdf:
        for i, page in enumerate(pdf.pages):
            text = page.extract_text() or ""
            if len(text.strip()) < 50:
                needs_ocr = True
            full_text_parts.append(text)

    full_text = "\n\n".join(full_text_parts).strip()

    if needs_ocr and len(full_text) < 200:
        if OCR_ENABLED:
            logger.info("PDF seems scanned, running OCR on %s", p.name)
            ocr_text = ocr_pdf(str(p))
            if ocr_text:
                full_text = ocr_text
            else:
                logger.warning("OCR returned empty for %s", p.name)
        else:
            logger.warning("PDF %s appears scanned but OCR is disabled", p.name)

    if full_text:
        docs.append({
            "title": p.stem,
            "text": full_text,
            "date": None,
            "meta": {"pages": len(full_text_parts), "format": "pdf"},
        })

    return docs


def _parse_docx(p: Path) -> list[dict]:
    """DOCX: параграфы + таблицы."""
    from docx import Document

    doc = Document(str(p))
    parts: list[str] = []

    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(", ".join(cells))

    text = "\n\n".join(parts)
    if text:
        return [{"title": p.stem, "text": text, "date": None, "meta": {"format": "docx"}}]
    return []


def _parse_doc(p: Path) -> list[dict]:
    """DOC: конвертация в DOCX через LibreOffice."""
    if not shutil.which("libreoffice") and not shutil.which("soffice"):
        raise RuntimeError(
            "LibreOffice not found. Install it to parse .doc files: "
            "apt install libreoffice-core libreoffice-writer"
        )

    import tempfile
    with tempfile.TemporaryDirectory() as tmpdir:
        cmd = [
            "libreoffice", "--headless", "--convert-to", "docx",
            "--outdir", tmpdir, str(p),
        ]
        result = subprocess.run(cmd, capture_output=True, text=True, timeout=60)
        if result.returncode != 0:
            raise RuntimeError(f"LibreOffice conversion failed: {result.stderr}")

        docx_path = Path(tmpdir) / (p.stem + ".docx")
        if not docx_path.exists():
            raise RuntimeError("Converted file not found after LibreOffice")

        return _parse_docx(docx_path)


def _parse_txt(p: Path) -> list[dict]:
    """TXT: автоопределение кодировки."""
    raw = p.read_bytes()
    detected = chardet.detect(raw)
    encoding = detected.get("encoding", "utf-8") or "utf-8"

    # Fallback
    try:
        text = raw.decode(encoding)
    except (UnicodeDecodeError, LookupError):
        try:
            text = raw.decode("cp1251")
        except UnicodeDecodeError:
            text = raw.decode("utf-8", errors="replace")

    text = text.strip()
    if text:
        return [{"title": p.stem, "text": text, "date": None, "meta": {"format": "txt", "encoding": encoding}}]
    return []


def _parse_md(p: Path) -> list[dict]:
    """Markdown: извлечение заголовков как title, остальное — текст."""
    text = p.read_text(encoding="utf-8", errors="replace")
    title = p.stem

    # Ищем первый заголовок
    match = re.search(r'^#\s+(.+)$', text, re.MULTILINE)
    if match:
        title = match.group(1).strip()

    text = text.strip()
    if text:
        return [{"title": title, "text": text, "date": None, "meta": {"format": "md"}}]
    return []


def _parse_json(p: Path) -> list[dict]:
    """JSON: рекурсивный обход, сбор строк > 50 символов."""
    data = json.loads(p.read_text(encoding="utf-8"))
    docs: list[dict] = []

    def _extract_strings(obj, path=""):
        if isinstance(obj, str) and len(obj) > 50:
            docs.append({
                "title": f"{p.stem} [{path}]" if path else p.stem,
                "text": obj,
                "date": None,
                "meta": {"format": "json", "path": path},
            })
        elif isinstance(obj, dict):
            for k, v in obj.items():
                _extract_strings(v, f"{path}.{k}" if path else k)
        elif isinstance(obj, list):
            for i, item in enumerate(obj):
                _extract_strings(item, f"{path}[{i}]")

    _extract_strings(data)
    return docs


def _parse_pptx(p: Path) -> list[dict]:
    """PPTX: слайды + заметки, каждый слайд = документ."""
    from pptx import Presentation

    prs = Presentation(str(p))
    docs: list[dict] = []

    for i, slide in enumerate(prs.slides):
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    if para.text.strip():
                        parts.append(para.text.strip())

        # Заметки
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"[Заметки]: {notes}")

        text = "\n".join(parts)
        if text:
            docs.append({
                "title": f"{p.stem} — Слайд {i + 1}",
                "text": text,
                "date": None,
                "meta": {"format": "pptx", "slide": i + 1},
            })

    return docs


def _parse_xlsx(p: Path) -> list[dict]:
    """XLSX: каждый лист = документ."""
    from openpyxl import load_workbook

    wb = load_workbook(str(p), read_only=True, data_only=True)
    docs: list[dict] = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows_text: list[str] = []
        headers: list[str] = []

        for i, row in enumerate(ws.iter_rows(values_only=True)):
            cells = [str(c).strip() for c in row if c is not None]
            if not cells:
                continue
            if i == 0:
                headers = cells
                continue
            if headers:
                pairs = [f"{h}: {v}" for h, v in zip(headers, cells) if v]
                rows_text.append(", ".join(pairs))
            else:
                rows_text.append(", ".join(cells))

        text = "\n".join(rows_text)
        if text:
            docs.append({
                "title": f"{p.stem} — {sheet_name}",
                "text": text,
                "date": None,
                "meta": {"format": "xlsx", "sheet": sheet_name},
            })

    wb.close()
    return docs


def _parse_csv(p: Path) -> list[dict]:
    """CSV: строки в текст."""
    raw = p.read_bytes()
    detected = chardet.detect(raw)
    encoding = detected.get("encoding", "utf-8") or "utf-8"

    try:
        text = raw.decode(encoding)
    except (UnicodeDecodeError, LookupError):
        text = raw.decode("utf-8", errors="replace")

    reader = csv.reader(io.StringIO(text))
    rows: list[str] = []
    headers: list[str] = []

    for i, row in enumerate(reader):
        if not row:
            continue
        if i == 0:
            headers = [c.strip() for c in row]
            continue
        pairs = [f"{h}: {c.strip()}" for h, c in zip(headers, row) if c.strip()]
        rows.append(", ".join(pairs))

    result = "\n".join(rows)
    if result:
        return [{"title": p.stem, "text": result, "date": None, "meta": {"format": "csv"}}]
    return []


def _parse_html(p: Path) -> list[dict]:
    """HTML: trafilatura для извлечения основного текста."""
    import trafilatura

    html = p.read_text(encoding="utf-8", errors="replace")
    text = trafilatura.extract(html, include_comments=False, include_tables=True)

    if not text or len(text.strip()) < 50:
        # Fallback: просто strips HTML tags
        text = re.sub(r'<[^>]+>', ' ', html)
        text = re.sub(r'\s+', ' ', text).strip()

    if text:
        return [{"title": p.stem, "text": text, "date": None, "meta": {"format": "html"}}]
    return []


def _parse_rtf(p: Path) -> list[dict]:
    """RTF: striprtf."""
    from striprtf.striprtf import rtf_to_text

    raw = p.read_bytes()
    detected = chardet.detect(raw)
    encoding = detected.get("encoding", "utf-8") or "utf-8"

    try:
        rtf_str = raw.decode(encoding)
    except (UnicodeDecodeError, LookupError):
        rtf_str = raw.decode("utf-8", errors="replace")

    text = rtf_to_text(rtf_str).strip()
    if text:
        return [{"title": p.stem, "text": text, "date": None, "meta": {"format": "rtf"}}]
    return []


def _parse_odt(p: Path) -> list[dict]:
    """ODT: конвертация в DOCX через LibreOffice, затем парсинг."""
    if not shutil.which("libreoffice") and not shutil.which("soffice"):
        raise RuntimeError(
            "LibreOffice not found. Install it to parse .odt files."
        )

    import tempfile
    with tempfile.TemporaryDirectory() as tmpdir:
        cmd = [
            "libreoffice", "--headless", "--convert-to", "docx",
            "--outdir", tmpdir, str(p),
        ]
        result = subprocess.run(cmd, capture_output=True, text=True, timeout=60)
        if result.returncode != 0:
            raise RuntimeError(f"LibreOffice ODT conversion failed: {result.stderr}")

        docx_path = Path(tmpdir) / (p.stem + ".docx")
        if not docx_path.exists():
            raise RuntimeError("Converted DOCX not found after LibreOffice")

        return _parse_docx(docx_path)


def _parse_epub(p: Path) -> list[dict]:
    """EPUB: каждая глава = документ."""
    from ebooklib import epub
    from bs4 import BeautifulSoup

    book = epub.read_epub(str(p))
    docs: list[dict] = []

    for item in book.get_items_of_type(9):  # ITEM_DOCUMENT
        html_content = item.get_content().decode("utf-8", errors="replace")
        soup = BeautifulSoup(html_content, "html.parser")
        text = soup.get_text(separator="\n").strip()

        if text and len(text) > 50:
            title = item.get_name() or p.stem
            docs.append({
                "title": f"{p.stem} — {title}",
                "text": text,
                "date": None,
                "meta": {"format": "epub", "chapter": title},
            })

    return docs


def _parse_image(p: Path) -> list[dict]:
    """Изображение: OCR."""
    if not OCR_ENABLED:
        logger.warning("OCR disabled, skipping image %s", p.name)
        return []

    text = ocr_image(str(p))
    if text:
        return [{"title": p.stem, "text": text, "date": None, "meta": {"format": "image"}}]
    return []
